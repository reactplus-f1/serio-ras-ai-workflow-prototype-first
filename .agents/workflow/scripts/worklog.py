"""Work log (作業ログ) in Word: one shared log for every AI/human activity.

Commands (run from repo root):
  python .agents/workflow/scripts/worklog.py template            # (re)build log/作業ログ_template.docx
  python .agents/workflow/scripts/worklog.py add --json entry.json   # append one activity (new L-xxx)
  python .agents/workflow/scripts/worklog.py set --id L-003 --field status --value "Thất bại"
  python .agents/workflow/scripts/worklog.py list               # print summary
  python .agents/workflow/scripts/worklog.py rebuild            # regenerate template + log in current format

Rules: .agents/workflow/rules/WORKLOG_RULES.md
"""
import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[3]
LOG_DIR = ROOT / "log"
LOG = LOG_DIR / "作業ログ.docx"
TEMPLATE = LOG_DIR / "作業ログ_template.docx"

# (key, label) — required items, in order.
FIELDS = [
    ("time", "Ngày và thời gian thực hiện"),
    ("run_id", "Run ID"),
    ("people", "Người thực hiện"),
    ("ticket", "Ticket / nội dung công việc"),
    ("type", "Loại hoạt động"),
    ("inputs", "Tài liệu / code đầu vào"),
    ("prompt", "Yêu cầu / chỉ thị gửi AI"),
    ("tool", "Công cụ AI và model"),
    ("commands", "Lệnh đã chạy và kết quả (exit code)"),
    ("ai_result", "Kết quả ban đầu do AI tạo"),
    ("human_changes", "Nội dung người đã hiệu chỉnh / bổ sung / loại bỏ"),
    ("review", "Kết quả review và test"),
    ("status", "Trạng thái"),
    ("time_ai", "Thời gian AI xử lý (phút)"),
    ("time_review", "Thời gian review thủ công (phút)"),
    ("time_rework", "Thời gian sửa lại (phút)"),
    ("notes", "Ghi chú / vấn đề phát sinh"),
]
LABEL = dict(FIELDS)
SUMMARY = ["ID", "Ngày giờ", "Run ID", "Người thực hiện", "Ticket / công việc", "Loại hoạt động", "Trạng thái",
           "AI (phút)", "Review (phút)", "Sửa lại (phút)"]
MINUTES = ("time_ai", "time_review", "time_rework")
TYPES = ["Chạy bước (AI tạo mới)", "Hiệu chỉnh", "Tạo lại", "Thử thất bại", "Chỉnh workflow",
         "Chỉnh repo code / môi trường"]
STATUSES = ["Thành công", "Thất bại", "Cần thực hiện lại"]
RULES = [
    "Mỗi dòng log tương ứng với một hoạt động tạo thay đổi (chạy bước, hiệu chỉnh, tạo lại, thử thất bại, chỉnh workflow/quy trình/skill, chỉnh repo code/môi trường), không tổng hợp một lần vào cuối ngày.",
    "Không ghi hỏi đáp, giải thích, trao đổi trong chat không làm thay đổi file/môi trường; kết quả review/test ghi vào mục của hoạt động tương ứng.",
    "Mục Yêu cầu / chỉ thị viết dạng chỉ thị chuyên môn (hành động — đối tượng — phạm vi/ràng buộc — tiêu chí), không chép nguyên văn lời chat.",
    "Ghi lại bản đầu tiên do AI tạo.",
    "Ghi lại các lần hiệu chỉnh, tạo lại và thử thất bại — mỗi lần một mục, không xóa mục cũ.",
    "Ghi rõ người thực hiện cho từng hoạt động.",
    "Run ID: <Function-ID>-S<n>-<yyyymmdd-HHMM> cho bước của run (ví dụ BR_Com_004-S3-20261017-0040); WF-<yyyymmdd-HHMM> cho chỉnh workflow/quy trình; chạy lại cùng bước thì Run ID mới.",
    "Lệnh đã chạy: mỗi lệnh build/test/kit/script kèm thư mục chạy, exit code và kết quả rút gọn để người khác chạy lại được.",
    "Thời gian tính bằng phút, tách riêng: AI xử lý (từ lúc gửi chỉ thị đến khi AI bàn giao), review thủ công (người đọc/kiểm tra kết quả), sửa lại (hiệu chỉnh sau review, cả người và AI). Không cộng trùng cùng một khoảng thời gian; thời gian chờ AI tính vào AI xử lý.",
    "Thời gian review/sửa lại do người phụ trách cung cấp; chưa có thì ghi “chờ user ghi” và cập nhật sau (worklog.py set). Mục cũ không đo được ghi “— (chưa đo)”.",
    "Không ghi secrets, mật khẩu, chuỗi kết nối.",
]


def _fonts(doc):
    st = doc.styles["Normal"]
    st.font.name = "Segoe UI"  # has Vietnamese glyphs; Japanese falls back to eastAsia font
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), "Segoe UI")
    rf.set(qn("w:eastAsia"), "Yu Gothic UI")
    st.font.size = Pt(9.5)


def _shade(cell, hex_="D9E2F3"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    cell._tc.get_or_add_tcPr().append(shd)


def _set(cell, text, bold=False):
    cell.text = ""
    r = cell.paragraphs[0].add_run(str(text)); r.bold = bold


def _detail(doc, lid, values):
    doc.add_heading(f"{lid} — {values.get('ticket', '')}", 2)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for key, label in FIELDS:
        row = t.add_row().cells
        _set(row[0], label, bold=True); _shade(row[0], "F2F2F2")
        _set(row[1], values.get(key, ""))
        row[0].width, row[1].width = Cm(5), Cm(12)
    doc.add_paragraph()


def build(path, sample):
    doc = Document()
    _fonts(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(1.5)
    doc.add_heading("作業ログ – Nhật ký công việc", 0)
    doc.add_paragraph("Một file log chung cho mọi hoạt động của workflow (mọi ticket/chức năng và chỉnh "
                      "workflow). Agent ghi bằng worklog.py ngay sau mỗi hoạt động; người phụ trách có thể "
                      "sửa trực tiếp trong Word (ví dụ ghi kết quả review/test).")
    doc.add_heading("1. Nội dung cần ghi", 1)
    for _, label in FIELDS:
        doc.add_paragraph(label, style="List Number")
    doc.add_paragraph("Loại hoạt động: " + ", ".join(TYPES) + ". Trạng thái: " + ", ".join(STATUSES) + ".")
    doc.add_heading("2. Quy tắc ghi log", 1)
    for r in RULES:
        doc.add_paragraph(r, style="List Bullet")
    doc.add_heading("3. Tổng hợp", 1)
    t = doc.add_table(rows=1, cols=len(SUMMARY)); t.style = "Table Grid"
    for i, h in enumerate(SUMMARY):
        _set(t.rows[0].cells[i], h, bold=True); _shade(t.rows[0].cells[i])
    doc.add_paragraph()
    doc.add_heading("4. Chi tiết", 1)
    if sample:
        _detail(doc, "L-xxx (mẫu)", {"ticket": "<Ticket / công việc>"})
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _summary_table(doc):
    for t in doc.tables:
        if [c.text for c in t.rows[0].cells][:2] == SUMMARY[:2]:
            return t
    raise SystemExit("Summary table not found in log")


def _next_id(table):
    nums = [int(r.cells[0].text[2:]) for r in table.rows[1:] if r.cells[0].text.startswith("L-")]
    return f"L-{(max(nums) + 1 if nums else 1):03d}"


def _summary_row(lid, v):
    return [lid, v.get("time", "")[:16], v.get("run_id", ""), v.get("people_short") or v.get("people", ""),
            v.get("summary") or v.get("ticket", ""), v.get("type", ""), v.get("status", "")] + [
            v.get(k, "") for k in MINUTES]


def _minutes(value):
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def add(values):
    missing = [k for k, _ in FIELDS if not str(values.get(k, "")).strip()]
    if missing:
        raise SystemExit("Missing fields: " + ", ".join(missing) + " (ghi “—” nếu không có)")
    if values["status"] not in STATUSES and not any(values["status"].startswith(s) for s in STATUSES):
        raise SystemExit("status must start with: " + ", ".join(STATUSES))
    for k in MINUTES:
        if _minutes(values[k]) is None and not values[k].startswith(("—", "chờ user ghi")):
            raise SystemExit(f"{k} must be minutes (number), “chờ user ghi” or “— (lý do)”")
    if not LOG.exists():
        build(LOG, sample=False)
    doc = Document(LOG)
    t = _summary_table(doc)
    lid = _next_id(t)
    cells = t.add_row().cells
    for i, v in enumerate(_summary_row(lid, values)):
        _set(cells[i], v)
    _detail(doc, lid, values)
    doc.save(LOG)
    print(f"{lid} added to {LOG.relative_to(ROOT).as_posix()}")


def set_field(lid, key, value):
    if key not in LABEL:
        raise SystemExit("field must be one of: " + ", ".join(LABEL))
    doc = Document(LOG)
    col = {k: i for i, k in enumerate(["id", "time", "run_id", "people", "ticket", "type", "status"] + list(MINUTES))}
    col.pop("id")
    if key in col:
        for r in _summary_table(doc).rows[1:]:
            if r.cells[0].text == lid:
                _set(r.cells[col[key]], value[:16] if key == "time" else value)
    body = list(doc.element.body)
    for i, el in enumerate(body):
        if el.tag == qn("w:p") and "".join(x.text or "" for x in el.iter(qn("w:t"))).startswith(lid + " "):
            tbl = next(e for e in body[i + 1:] if e.tag == qn("w:tbl"))
            for tr in tbl.iter(qn("w:tr")):
                tcs = tr.findall(qn("w:tc"))
                if "".join(x.text or "" for x in tcs[0].iter(qn("w:t"))) == LABEL[key]:
                    from docx.table import _Cell
                    _set(_Cell(tcs[1], None), value)
            doc.save(LOG)
            print(f"{lid}.{key} = {value}")
            return
    raise SystemExit(f"{lid} not found")


def rows(path=LOG, header=False):
    t = _summary_table(Document(path))
    out = [[c.text for c in r.cells] for r in t.rows]
    return out if header else out[1:]


def entries(path=LOG):
    """Read every detail table back as {key: value}; labels no longer in FIELDS are dropped."""
    doc = Document(path)
    key_of = {label: key for key, label in FIELDS}
    key_of["Yêu cầu / prompt gửi AI"] = "prompt"  # old label
    head, *body_rows = rows(path, header=True)
    summ = {r[0]: dict(zip(head, r)) for r in body_rows}
    body, out = list(doc.element.body), []
    for i, el in enumerate(body):
        txt = "".join(x.text or "" for x in el.iter(qn("w:t"))) if el.tag == qn("w:p") else ""
        lid = txt.split(" ")[0] if txt else ""
        if lid in summ and " — " in txt:
            tbl = next(e for e in body[i + 1:] if e.tag == qn("w:tbl"))
            v = {}
            for tr in tbl.iter(qn("w:tr")):
                tcs = tr.findall(qn("w:tc"))
                label = "".join(x.text or "" for x in tcs[0].iter(qn("w:t")))
                if label in key_of:
                    v[key_of[label]] = "\n".join("".join(x.text or "" for x in p.iter(qn("w:t")))
                                                 for p in tcs[1].iter(qn("w:p")))
            v["people_short"] = summ[lid].get("Người thực hiện", "")
            v["summary"] = summ[lid].get("Ticket / công việc", "")
            out.append(v)
    return out


def rebuild():
    """Regenerate template and log in the current format, keeping every entry and its order."""
    old = entries() if LOG.exists() else []
    backup = LOG.with_name(LOG.stem + "_backup.docx")
    if LOG.exists():
        backup.write_bytes(LOG.read_bytes())
    build(TEMPLATE, sample=True)
    build(LOG, sample=False)
    legacy = {"run_id": "— (mục cũ, trước khi áp dụng Run ID)",
              "commands": "— (mục cũ, không ghi lệnh)"}
    for v in old:
        for key, _ in FIELDS:
            v.setdefault(key, legacy.get(key, "— (chưa đo)" if key in MINUTES else "—"))
        add(v)
    backup.unlink(missing_ok=True)


def main():
    for stream in (sys.stdout, sys.stderr, sys.stdin):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="backslashreplace")
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("template")
    a = sub.add_parser("add"); a.add_argument("--json", required=True, help="file JSON, '-' = stdin")
    s = sub.add_parser("set"); s.add_argument("--id", required=True); s.add_argument("--field", required=True); s.add_argument("--value", required=True)
    sub.add_parser("list")
    sub.add_parser("rebuild")
    args = ap.parse_args()
    if args.cmd == "template":
        build(TEMPLATE, sample=True); print(TEMPLATE.relative_to(ROOT).as_posix())
    elif args.cmd == "add":
        raw = sys.stdin.read() if args.json == "-" else Path(args.json).read_text(encoding="utf-8")
        data = json.loads(raw)
        for v in (data if isinstance(data, list) else [data]):
            add(v)
    elif args.cmd == "rebuild":
        rebuild()
    elif args.cmd == "set":
        set_field(args.id, args.field, args.value)
    else:
        rs = rows()
        for r in rs:
            print(" | ".join(r))
        totals = [sum(m for m in (_minutes(r[SUMMARY.index(h)]) for r in rs) if m is not None)
                  for h in SUMMARY[-3:]]
        print(f"{len(rs)} entries | tổng phút AI {totals[0]:g}, review {totals[1]:g}, sửa lại {totals[2]:g} (mục chưa đo không tính)")


if __name__ == "__main__":
    try:
        main()
    except PermissionError as exc:
        raise SystemExit(f"Không ghi được {exc.filename}: file đang mở trong Word — đóng file rồi chạy lại.")
