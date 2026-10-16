from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import sys

# Báo cáo tiến độ duy nhất của màn hình: cập nhật dồn sau mỗi bước, chạy lại để ghi đè .docx (WORKFLOW.md "Báo cáo tiến độ").
OUT = sys.argv[1] if len(sys.argv) > 1 else __file__.replace("make_report.py", "報告_BR_Com_004.docx")
RUN = "serio-ai-workflow/runs/BR_Com_004"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Yu Gothic UI"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic UI")
st.font.size = Pt(10.5)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin = s.bottom_margin = Cm(1.8)


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h_ in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h_)
        r.bold = True
        shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p_ = cells[i].paragraphs[0]
            if isinstance(v, tuple):
                r = p_.add_run(v[0]); r.bold = True
                r.font.color.rgb = RGBColor.from_string(v[1])
            else:
                p_.add_run(str(v))
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def h(text, lvl=1):
    doc.add_heading(text, lvl)


def p(text, bold=False):
    para = doc.add_paragraph()
    r = para.add_run(text); r.bold = bold
    return para


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


OK = ("Đạt", "1E7B34")
PART = ("Một phần", "B7791F")

# ---------------- Title ----------------
doc.add_heading("Báo cáo thực hiện – BR_Com_004 会議室タブレット登録 (cập nhật đến S2, chờ review G1)", 0)
p("PA1: 方式1 – Tạo thành phẩm đầu tiên bằng AI (thay cho DD: thành phẩm cuối bước 1 = Hướng dẫn triển khai prototype pg-guide.md; bước 2 = code PG trong serio-ras)")

table(["Mục", "Nội dung"], [
    ["Ticket", "PA1 – 方式1：BR004 (thời gian dự kiến 2026/10/15 – 2026/10/16)"],
    ["Chức năng", "BR_Com_004 会議室タブレット登録 (BR_Com_003 / BR_Com_005 là mock)"],
    ["Bước / trạng thái", "S1 hoàn thành → G0 xác nhận trong chat (user cung cấp API一覧, テーブル設計書 và chỉ thị chạy S2) → S2 hoàn thành → đang dừng tại điểm review G1 (user review code PG)"],
    ["Ngày thực hiện", "S1: 2026/10/15 17:58 – 18:15; bổ sung QA.md 23:57. S2: 2026/10/16 00:00 – 00:58 (giờ VN)"],
    ["Người thực hiện", "AI: GitHub Copilot CLI (model Claude Fable 5.1) – thao tác: PC"],
    ["Kết quả checker", "kit.py check → S1: READY, S2: READY; kit.py next → review point G1, bước kế tiếp S3 (Testspec)"],
    ["Thư mục hồ sơ", RUN + " (hồ sơ); D:/Serio-AI/serio-ras branch bd-to-pg (code)"],
], widths=[4, 13])

# ---------------- Checklist ----------------
h("1. Checklist công việc (theo ticket, DD thay bằng thành phẩm cuối bước 1)")
table(["#", "Hạng mục", "Trạng thái", "Đã làm / bằng chứng"], [
    ["1", "Chuẩn bị đầy đủ tài liệu và source đầu vào cho AI", OK,
     "S1: BD 【RAS】(BR_Com_004)会議室タブレット登録_画面定義書.xlsx Ver 1.01, scope chốt trong chat, bộ rule .agents/workflow/rules, repo serio-ras. "
     "Trước S2 user bổ sung API一覧.xlsx, テーブル設計書.xlsx (đọc trong chat, không lưu vào runs) và dependency-analysis.md. "
     "Chưa có: 画面共通仕様説明書, メッセージ一覧, BD của BR_Com_003/005 → ghi PENDING; ảnh hưởng: R-026 chuỗi cấm chưa implement, ID message tạm."],
    ["2", "Sử dụng AI để tạo bản đầu tiên cho toàn bộ BR004", OK,
     "S1: pg-guide.md bao phủ 4 event (初期表示, chọn 設置会議室, search 延長候補会議室, 登録): 28 Rule (R-001…R-028), 15 atomic condition (C-001…C-015), 18 giả định (A-01…A-18), 17 mục kế hoạch code (P-01…P-17); QA.md 14 câu hỏi nghiệp vụ. "
     "S2: code 3 API (API_004/005/008), entity/repository/usecase/controller backend, màn Angular BRCom004 + 2 mock; dotnet build 0 error, ng build (development) 0 error."],
    ["3", "Tạo thành phẩm theo format nội bộ đã thống nhất", OK,
     "Markdown theo template .agents/workflow/templates (pg-guide, task, source-manifest, QA, pg-implementation, pg-artifact-manifest, process-log). "
     "Checker kit.py xác nhận cấu trúc Rules / Atomic conditions / Kế hoạch code (S1) và manifest file tồn tại (S2)."],
    ["4", "Lưu lại nguyên trạng bản đầu tiên do AI tạo", OK,
     "Bản AI tạo giữ nguyên trong runs/BR_Com_004/ và serio-ras (chưa ai sửa nội dung thiết kế/code). "
     "Điểm review G0–G5 do user xác nhận trong chat; user tự commit (không dùng approvals.md, không khóa bằng SHA)."],
    ["5", "Ghi lại prompt, công cụ AI, kết quả và công số (#9)", OK,
     "runs/BR_Com_004/process/process-log.md có 3 dòng (S1, QA/skill update, S2): prompt/procedure, tool (Copilot CLI, kit 0.3.0, xlsx_view.py, dotnet 10.0.302, node 24.3.0), model, kết quả check/next, giới hạn. Công số: xem mục 5."],
    ["6", "Ghi rõ nguồn tham chiếu cho từng nội dung thiết kế", OK,
     "pg-guide.md: mọi dòng có [SRC: BD004 | sheet | ô] hoặc [DECISION: A-xx]. pg-implementation.md: mỗi module ghi Rule IDs và Source/decision; code có comment Rule ID / SELECT_xxx / INSERT_xxx tại chỗ xử lý."],
    ["7", "Không tự bổ sung yêu cầu khi không có căn cứ", OK,
     "Chỗ BD không nói tới được ghi thành giả định A-01…A-18 (S1) và lệch/giả định mới D-1…D-8, A-19 (S2) với lý do và ảnh hưởng, để user chấp nhận/sửa tại G1. "
     "Rule chỉ dựa trên giả định đánh Required = NO (R-010, R-013, R-026, R-028)."],
], widths=[0.8, 4.2, 2, 10])

# ---------------- Inputs ----------------
h("2. Đầu vào đã dùng")
table(["ID", "Tài liệu", "Phiên bản", "Vai trò", "Trạng thái"], [
    ["BD01", "【RAS】(BR_Com_004)会議室タブレット登録_画面定義書.xlsx", "Ver 1.01 (2026/07/31, SRO 牟禮)", "BD chính", "Đã đọc đủ 14 sheet (S1)"],
    ["—", "【RAS】API一覧.xlsx", "user đính kèm chat 2026/10/15", "Endpoint, role API_004/005/008", "Đã đọc (S2); không lưu vào runs theo yêu cầu user"],
    ["—", "【RAS】テーブル設計書.xlsx", "user đính kèm chat 2026/10/15", "Schema rooms_mst / tablets_mst / candidate_rooms_mst", "Đã đọc (S2); khớp ras_tables.sql"],
    ["—", "runs/BR_Com_004/dependency-analysis.md", "user tạo 2026/10/16 00:22", "Phân tích phụ thuộc, contract mock D-01/D-02", "Đã dùng (S2)"],
    ["—", "Scope (chat 2026/10/15)", "—", "Chức năng chính BR_Com_004; BR_Com_003/005 mock", "Đã ghi task.md"],
    ["—", ".agents/workflow/rules (PROJECT, SOURCE_MAP, UI, ARCHITECTURE, REPOSITORY, CODING, BR_COM_004)", "kit 0.3.0", "Rule tổng hợp", "Đã đọc"],
    ["—", "Source serio-ras (ras-system, ras-web)", "branch bd-to-pg", "Pattern code, shared components", "Đã đọc trước khi code (S2)"],
    ["—", "画面共通仕様説明書, メッセージ一覧", "—", "Chuỗi cấm, ID message", "PENDING"],
    ["—", "BD BR_Com_003 / BR_Com_005", "—", "Mock contract", "PENDING (mock theo A-07/A-08)"],
], widths=[1.2, 5.5, 4, 3.5, 2.8])
p("Giới hạn đọc BD: sheet レイアウト là ảnh/shape, script không đọc được → cần người xem Excel đối chiếu bố cục.")

# ---------------- Outputs S1 ----------------
h("3. Thành phẩm bước 1 (S1)")
table(["File", "Nội dung"], [
    [RUN + "/pg-guide.md", "Hướng dẫn triển khai prototype: BD & phạm vi, sheet đã đọc, luồng event, field/state, nhánh lỗi, dữ liệu/API, mock, giả định A-01…A-18, Rules, Atomic conditions, Kế hoạch code P-01…P-17"],
    [RUN + "/task.md", "Function ID, BD path, scope, ngoài scope, repo (serio-ras), dependency thật/mock"],
    [RUN + "/source-manifest.md", "Nguồn thật đã dùng (ID / Path / Role)"],
    [RUN + "/QA.md", "14 câu hỏi nghiệp vụ gửi Serio (No. / 該当箇所 / 質問 / 回答), tiếng Nhật"],
    [RUN + "/process/process-log.md", "Process log S1, QA, S2"],
    [RUN + "/inputs/<BD>.xlsx", "Bản sao BD làm việc"],
], widths=[6.5, 10.5])

# ---------------- Outputs S2 ----------------
h("4. Thành phẩm bước 2 (S2 – code PG)")
p("Hồ sơ:", bold=True)
table(["File", "Nội dung"], [
    [RUN + "/pg-implementation.md", "Changes P-01…P-17 (file, Rule IDs, source, impact), lệch/giả định mới D-1…D-8 + A-19, inventory branch/condition C-001…C-015 cho S3, bằng chứng build/lint"],
    [RUN + "/pg-artifact-manifest.md", "37 file (28 mới, 7 sửa, 2 build-config) – ID / Path / Role"],
], widths=[6.5, 10.5])

p("Code trong serio-ras (branch bd-to-pg, user tự commit):", bold=True)
table(["Lớp", "File chính", "Nội dung"], [
    ["RasDomain", "Models/Room.cs, Tablet.cs, CandidateRoom.cs; Repositories/IRoomRepository.cs, ITabletRepository.cs", "Entity 3 bảng theo テーブル設計書; interface SELECT_001…004, INSERT/UPDATE/DELETE"],
    ["RasInfrastructure", "RasDbContext.cs (sửa); EntityConfigurations/*3; Repositories/RoomRepository.cs, TabletRepository.cs", "Mapping EF Core (UNIQUE machine_id, PK ghép), truy vấn SELECT_002 outer join, check No.3, replace candidate"],
    ["RasUseCase", "ViewModels/TabletInfoModels.cs; Services/ITabletInfoService.cs, TabletInfoService.cs", "Nghiệp vụ 登録: trim ID, check trùng → 422, INSERT_001/UPDATE_001, DELETE_001+INSERT_002, 1 transaction"],
    ["RasApi", "Controllers/TabletInfoController.cs (API_004/005, AdministratorOnly), RoomsController.cs (API_008); ViewModels/TabletInfoRequests.cs; Startup.cs (DI, sửa)", "Endpoint theo API一覧, validation 400, envelope ApiResponse, 422 qua CreateFailResponse"],
    ["Common", "Util/Resources/Messages.resx (sửa)", "E_BRCOM004_001 = タブレット端末ID：他の会議室に登録されています。（会議室名：{0}）"],
    ["ras-web feature", "features/tablet-registration/** (component, service, interface)", "Màn BRCom004: dropdown 設置会議室, タブレット端末ID, 延長候補会議室 (join ;), 登録 với confirm/success/422"],
    ["ras-web mock", "features/room-search-mock/**, features/error-mock/**", "Mock BR_Com_003 (modal chọn phòng) và BR_Com_005 (route error/:statusCode)"],
    ["ras-web shared (sửa)", "app.routes.ts, guards/admin.guard.ts, services/api-error-handler.service.ts, assets/i18n/ja.json", "Route BRCom004 (AuthGuard+AdminGuard), lỗi quyền/API → /error/{status}, message/label tiếng Nhật"],
], widths=[3, 7, 7])

p("Bằng chứng build/lint:", bold=True)
bullets([
    "Backend: dotnet build src/Services/Ras/RasApi/RasApi.csproj → Build succeeded, 0 error; 19 warning đều có sẵn trước thay đổi.",
    "Frontend: ng build --configuration development → 0 error. npm run build (production) compile OK nhưng vượt budget 1.20 MB có sẵn (1.47 MB); npm run lint lỗi cấu hình eslint có sẵn (rule prefer-style-binding không tồn tại) – không do S2.",
])

# ---------------- Effort ----------------
h("5. Prompt, công cụ AI và công số (#9)")
table(["Mục", "Giá trị"], [
    ["Công cụ / model", "GitHub Copilot CLI (app Copilot), model Claude Fable 5.1; kit.py 0.3.0; xlsx_view.py; dotnet SDK 10.0.302; node 24.3.0 / Angular CLI"],
    ["Prompt S1", "“Đọc AGENTS.md và .agents/workflow/START_CHAT.md trong serio-ai-workflow, rồi chạy S1 cho file BD đính kèm. Scope: chỉ chức năng chính của BR_Com_004; BR_Com_003/005 là mock.”"],
    ["Prompt S2", "“từ kết quả chạy S1 hãy thực hiện S2” → đính kèm API一覧, テーブル設計書 → “tiến hành chạy S2, không cần lưu file này vào runs”"],
    ["Procedure AI đã theo", "S1: AGENTS → START_CHAT → WORKFLOW → PROJECT/SOURCE_MAP → skill 01-pg-guide → UI/ARCHITECTURE/REPOSITORY/BR_COM_004 rules → templates. S2: skill 02-pg → CODING/ARCHITECTURE/REPOSITORY/UI rules → đọc source serio-ras → code theo P-01…P-17 → build → pg-implementation/manifest"],
    ["Lệnh kit đã chạy", "init → next (STEP S1) → check (S1 READY) → [G0 xác nhận chat] → next (STEP S2) → check (S2 READY) → next (review G1, STEP S3)"],
    ["Thời gian máy", "S1 ≈ 17 phút (1 lượt chat) + QA/skill ≈ 3 lượt; S2 ≈ 60 phút (1 lượt chat dài, gồm đọc source và build)"],
    ["Thời gian người", "Review G0 thực hiện gián tiếp (user cung cấp input và chỉ thị chạy S2); review G1 chưa thực hiện"],
    ["Sửa thủ công nội dung thiết kế/code", "0. Thay đổi workflow theo yêu cầu user trong quá trình: bỏ SHA-256 file hash; bỏ approvals.md (gate xác nhận trong chat); thêm QA.md là output S1; bỏ yêu cầu Code commit SHA (user tự commit, không khóa bước)"],
], widths=[4, 13])

# ---------------- Review G1 ----------------
h("6. Điểm cần user quyết định tại review G1")
table(["ID", "Vấn đề", "AI chọn", "Ảnh hưởng nếu đổi"], [
    ["D-1", "DELETE_001 + INSERT_002", "Replace theo diff (xóa dòng không còn, thêm dòng mới) thay vì xóa hết rồi insert", "Kết quả DB giống BD; update_at dòng giữ nguyên không đổi"],
    ["D-2", "Transaction R-023", "Một SaveChangesAsync (EF Core tự bọc transaction)", "Không có bước ghi ngoài SaveChanges → all-or-nothing"],
    ["D-3", "ID message", "E_BRCOM004_001 (backend), message.tabletRegistration.* (FE)", "Đổi khi có メッセージ一覧"],
    ["D-4", "R-026 chuỗi cấm (Required NO)", "Chưa implement", "Bổ sung khi có 画面共通仕様"],
    ["D-5", "Separator 延長候補会議室", "“;” không khoảng trắng (BD I10) dù Constants.separator = “; ”", "Bỏ .trim() nếu Serio chọn “; ”"],
    ["D-6", "API_004 khi không có tablet", "200 với data = null (FE clear)", "Nếu muốn 404 cần đổi FE handler"],
    ["D-8", "AdminGuard / ApiErrorHandler", "Mọi lỗi 400/401/403/404/409/500/504 → /error/{status} (mock BR_Com_005)", "Hành vi màn sample khi lỗi API thay đổi (trước: về /sample)"],
    ["A-19", "SELECT_001 với is_deleted", "Không lọc (BD không ghi điều kiện)", "Thêm Where(!IsDeleted) nếu cần"],
    ["A-02…A-18", "Giả định S1 còn hiệu lực", "Xem pg-guide.md mục “Giả định của AI”", "Trả lời QA.md khác giả định → sửa pg-guide, làm lại từ G1"],
], widths=[1.5, 4, 6.5, 5])

# ---------------- Next ----------------
h("7. Bước tiếp theo")
bullets([
    "User review code PG + pg-implementation.md, xác nhận G1 trong chat (hoặc yêu cầu sửa → AI làm lại từ G1).",
    "User tự commit: serio-ras (28 file mới, 7 file sửa – xem pg-artifact-manifest.md) và serio-ai-workflow (runs/BR_Com_004/**, .agents/** thay đổi workflow). AI không thao tác Git.",
    "Kiểm tra chạy thử: tạo DB bằng ras-system/sql/ras_tables.sql + seed rooms_mst; dotnet run RasApi (https://localhost:5101); npm start ras-web (http://localhost:4200/BRCom004); tài khoản Entra ID có role Administrator. Kịch bản 11 bước đã gửi trong chat 2026/10/16.",
    "Cung cấp cho S3 (Testspec): 単体テスト仕様書_ルール (bắt buộc); nếu có: 画面共通仕様説明書, メッセージ一覧, câu trả lời QA.md từ Serio.",
])

doc.save(OUT)
print("saved", OUT.encode("ascii", "backslashreplace").decode())
